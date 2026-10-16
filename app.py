import io
import os
import zipfile
from pathlib import Path

import streamlit as st
import xlrd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="UNTAPPED STUDIO | 나염·발주 엑셀 자동 가공", page_icon="⚙️", layout="wide")

def clean_cell_value(val):
    """xlrd가 float로 읽는 정수형 데이터(바코드, 상품코드 등)를 정수로 정리."""
    if isinstance(val, float) and val.is_integer():
        return int(val)
    return val

def process_single_file(file_bytes, input_filename):
    """업로드된 .xls 파일을 기존 매크로 규칙으로 가공해 .xlsx 바이트를 반환."""
    try:
        source = io.BytesIO(file_bytes)
        source_wb = xlrd.open_workbook(file_contents=source.getvalue())
        sheet = source_wb.sheet_by_index(0)

        top_title = sheet.cell_value(0, 0)

        # 기존 코드와 동일하게 11번째 행부터 마지막 행 직전까지 데이터를 읽습니다.
        rows_data = []
        for r in range(10, sheet.nrows - 1):
            c0 = clean_cell_value(sheet.cell_value(r, 0))
            c2 = clean_cell_value(sheet.cell_value(r, 2))
            c3 = clean_cell_value(sheet.cell_value(r, 3))
            c5 = clean_cell_value(sheet.cell_value(r, 5))
            c4 = clean_cell_value(sheet.cell_value(r, 4))
            rows_data.append([c0, c2, c3, c5, c4])

        grouped_rows = []
        prev_item_code = None
        prev_logo_part = None

        for row in rows_data:
            item_code = str(row[1])
            opt_str = str(row[2])
            logo_part = opt_str.rsplit("-", 1)[0] if "-" in opt_str else opt_str

            if prev_item_code is not None:
                if item_code != prev_item_code:
                    grouped_rows.append((["", "", "", "", ""], "item_change"))
                elif logo_part != prev_logo_part:
                    grouped_rows.append((["", "", "", "", ""], "logo_change"))

            grouped_rows.append((row, "data"))
            prev_item_code = item_code
            prev_logo_part = logo_part

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "가공리스트"
        ws.views.sheetView[0].showGridLines = True

        thin = Side(border_style="thin", color="000000")
        border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

        ws.merge_cells("A1:E2")
        cell_a1 = ws["A1"]
        cell_a1.value = top_title
        cell_a1.font = Font(name="Malgun Gothic", size=16, bold=True)
        cell_a1.alignment = Alignment(horizontal="center", vertical="center")

        for r in range(1, 3):
            ws.row_dimensions[r].height = 25
            for c in range(1, 6):
                ws.cell(row=r, column=c).border = border_all

        headers = ["공급처상품명", "상품코드", "옵션", "요청수량", "바코드"]
        header_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        ws.row_dimensions[3].height = 30

        for c_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=c_idx, value=header)
            cell.font = Font(name="Malgun Gothic", size=13, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.fill = header_fill
            cell.border = border_all

        item_change_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        curr_row = 4
        total_qty = 0

        for r_data, r_type in grouped_rows:
            ws.row_dimensions[curr_row].height = 30

            if r_type == "item_change":
                for c_idx in range(1, 6):
                    cell = ws.cell(row=curr_row, column=c_idx, value="")
                    cell.border = border_all
                    cell.fill = item_change_fill
            elif r_type == "logo_change":
                for c_idx in range(1, 6):
                    cell = ws.cell(row=curr_row, column=c_idx, value="")
                    cell.border = border_all
            else:
                for c_idx, val in enumerate(r_data, 1):
                    cell = ws.cell(row=curr_row, column=c_idx, value=val)
                    cell.font = Font(name="Malgun Gothic", size=13, bold=True)
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
                    cell.border = border_all

                    if c_idx == 4:
                        try:
                            qty_val = int(val)
                            total_qty += qty_val
                            cell.value = qty_val
                            cell.number_format = "#,##0"
                        except (ValueError, TypeError):
                            pass
                    elif c_idx in (2, 5) and isinstance(val, (int, float)):
                        cell.number_format = "@"
            curr_row += 1

        ws.row_dimensions[curr_row].height = 30
        cell_total = ws.cell(row=curr_row, column=1, value="합계")
        cell_total.font = Font(name="Malgun Gothic", size=13, bold=True)
        cell_total.alignment = Alignment(horizontal="center", vertical="center")
        cell_total.border = border_all

        for c_idx in range(2, 6):
            cell = ws.cell(row=curr_row, column=c_idx)
            cell.border = border_all
            if c_idx == 4:
                cell.value = total_qty
                cell.font = Font(name="Malgun Gothic", size=13, bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.number_format = "#,##0"

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row < 3:
                    continue
                if cell.value is not None:
                    val_str = str(cell.value)
                    length = sum(2 if ord(char) > 127 else 1 for char in val_str)
                    max_len = max(max_len, length)
            ws.column_dimensions[col_letter].width = max(max_len * 1.3 + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        base_name = Path(input_filename).stem
        output_name = f"[가공완료]_{base_name}.xlsx"
        return output_name, output.getvalue(), total_qty, len(rows_data)
    except Exception as exc:
        raise RuntimeError(f"{input_filename}: {exc}") from exc

st.markdown("""
<style>
.stApp {background: linear-gradient(135deg, #f5faff 0%, #ffffff 65%, #eef7ff 100%);}
.topbar {background:#164b91;color:white;padding:18px 24px;border-radius:14px;margin-bottom:16px;}
.topbar h1 {font-size:28px;margin:0;font-weight:800;}
.notice {background:#e8f3ff;border:1px solid #b9dcff;border-radius:10px;padding:13px 16px;color:#153b70;margin-bottom:18px;}
.panel {border:1px solid #b9d9f5;border-radius:14px;padding:18px;background:rgba(255,255,255,.9);}
</style>
<div class="topbar"><h1>⚙️ UNTAPPED STUDIO | 나염·발주 엑셀 자동 가공 웹매크로</h1></div>
<div class="notice">💡 파이썬 미설치 PC에서도 링크로 접속해 원본 .xls 파일을 업로드하면 가공된 .xlsx 파일로 변환할 수 있습니다.</div>
""", unsafe_allow_html=True)

left, right = st.columns([1, 1], gap="large")
with left:
    st.markdown("### 🖥️ 원본 엑셀 업로드")
    st.markdown('<div class="panel">여러 개의 원본 <b>.xls</b> 파일을 한 번에 선택하세요.</div>', unsafe_allow_html=True)
    uploaded_files = st.file_uploader(
        "원본 .xls 파일 선택 (여러 파일 선택 가능)",
        type=["xls"],
        accept_multiple_files=True,
        help="현재 버전은 기존 매크로와 동일하게 .xls 파일을 입력으로 받습니다."
    )
    st.caption("※ 파일은 가공을 위해 서버 메모리에서 처리되며, 이 앱은 업로드 파일을 별도 저장하지 않습니다.")

with right:
    st.markdown("### 📊 변환 결과")
    st.markdown('<div class="panel">상품코드별 구분 행, 로고·컬러 구분 행, 요청수량 합계 및 엑셀 서식을 자동 적용합니다.</div>', unsafe_allow_html=True)
    st.markdown("- 상품코드 변경: 연노란색 구분 행")
    st.markdown("- 동일 상품코드 내 로고·컬러 변경: 일반 구분 행")
    st.markdown("- 결과 시트명: `가공리스트`")
    st.markdown("- 파일별 다운로드 및 전체 ZIP 다운로드")

if uploaded_files:
    st.info(f"업로드된 파일: {len(uploaded_files)}개")
    if st.button("⚡ 엑셀 자동 가공 시작", type="primary", use_container_width=True):
        results = []
        errors = []
        progress = st.progress(0)
        status = st.empty()
        for idx, uploaded in enumerate(uploaded_files):
            status.write(f"가공 중: {uploaded.name}")
            try:
                results.append(process_single_file(uploaded.getvalue(), uploaded.name))
            except Exception as exc:
                errors.append(str(exc))
            progress.progress((idx + 1) / len(uploaded_files))
        status.empty()
        progress.empty()

        if results:
            st.success(f"가공 완료: {len(results)}개 파일")
            zip_buffer = io.BytesIO()
            with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
                for output_name, output_bytes, total_qty, row_count in results:
                    zf.writestr(output_name, output_bytes)
            st.download_button(
                "📦 전체 가공완료 파일 ZIP 다운로드",
                data=zip_buffer.getvalue(),
                file_name="나염리스트_가공완료_전체.zip",
                mime="application/zip",
                type="primary",
                use_container_width=True
            )
            for output_name, output_bytes, total_qty, row_count in results:
                c1, c2, c3 = st.columns([3, 1, 2])
                c1.write(f"📄 **{output_name}**")
                c2.write(f"수량 합계: **{total_qty:,}**")
                c3.download_button(
                    "⬇️ 개별 다운로드",
                    data=output_bytes,
                    file_name=output_name,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key=f"download_{output_name}"
                )
        if errors:
            st.error("일부 파일을 변환하지 못했습니다.")
            for err in errors:
                st.write(f"- {err}")

st.markdown("---")
st.caption("UNTAPPED STUDIO · 나염·발주 엑셀 자동 가공 웹매크로")
